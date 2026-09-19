import asyncio
import hashlib
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Protocol
from uuid import UUID, uuid4

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.core.config import Settings
from app.models.domain import (
    AuditEvent,
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeVersion,
)
from app.models.enums import DocumentProcessingStatus, KnowledgeStatus
from app.services.embeddings import EmbeddingProvider, EmbeddingProviderError

READ_BLOCK_BYTES = 1024 * 1024
EMBEDDING_BATCH_SIZE = 32
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 250
SUPPORTED_DOCUMENT_TYPES = frozenset(
    {
        "house_manual",
        "house_rules",
        "appliance_guide",
        "check_in_guide",
        "local_guide",
        "other",
    }
)

_MEDIA_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
_ACCEPTED_DECLARED_MEDIA_TYPES = {
    ".txt": {"text/plain", "application/octet-stream"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".html": {"text/html", "application/octet-stream"},
    ".htm": {"text/html", "application/octet-stream"},
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
        "application/zip",
    },
}


class DocumentUploadError(Exception):
    def __init__(self, status_code: int, detail: str) -> None:
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


class DocumentProcessingError(Exception):
    pass


class DocumentJobQueue(Protocol):
    async def process(self, document_id: UUID) -> None: ...


class UploadSource(Protocol):
    @property
    def filename(self) -> str | None: ...

    @property
    def content_type(self) -> str | None: ...

    async def read(self, size: int = -1) -> bytes: ...


@dataclass(frozen=True)
class StoredUpload:
    original_filename: str
    storage_key: str
    media_type: str
    size_bytes: int
    file_sha256: str


def sanitize_filename(filename: str | None) -> str:
    leaf = re.split(r"[/\\]", filename or "")[-1]
    leaf = "".join(character for character in leaf if character.isprintable())
    leaf = re.sub(r"\s+", " ", leaf).strip(" .")
    if not leaf:
        return "document"
    suffix = Path(leaf).suffix[:16]
    stem_limit = max(1, 255 - len(suffix))
    return f"{Path(leaf).stem[:stem_limit]}{suffix}"[:255]


def _extension(filename: str) -> str:
    extension = Path(filename).suffix.casefold()
    if extension not in _MEDIA_TYPES:
        supported = ", ".join(sorted(_MEDIA_TYPES))
        raise DocumentUploadError(415, f"Unsupported file type. Supported types: {supported}")
    return extension


def _validate_declared_media_type(extension: str, media_type: str | None) -> None:
    declared = (media_type or "application/octet-stream").split(";", maxsplit=1)[0].casefold()
    if declared not in _ACCEPTED_DECLARED_MEDIA_TYPES[extension]:
        raise DocumentUploadError(415, "The file extension and media type do not match")


def _validate_file_signature(extension: str, data: bytes) -> None:
    if not data:
        raise DocumentUploadError(400, "The uploaded file is empty")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise DocumentUploadError(400, "The uploaded PDF has an invalid signature")
    if extension == ".docx":
        if not data.startswith(b"PK"):
            raise DocumentUploadError(400, "The uploaded DOCX has an invalid signature")
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if "word/document.xml" not in {entry.filename for entry in entries}:
                    raise DocumentUploadError(400, "The uploaded file is not a valid DOCX")
                if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise DocumentUploadError(413, "The expanded DOCX is too large")
        except zipfile.BadZipFile as exc:
            raise DocumentUploadError(400, "The uploaded DOCX is invalid") from exc
    if extension in {".txt", ".md", ".html", ".htm"}:
        if b"\x00" in data:
            raise DocumentUploadError(400, "The uploaded text file contains binary data")
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentUploadError(400, "Text documents must use UTF-8 encoding") from exc


def resolve_storage_path(storage_root: Path, storage_key: str) -> Path:
    key = PurePosixPath(storage_key)
    if key.is_absolute() or ".." in key.parts:
        raise DocumentProcessingError("The stored document path is invalid")
    root = storage_root.resolve()
    resolved = root.joinpath(*key.parts).resolve()
    if not resolved.is_relative_to(root):
        raise DocumentProcessingError("The stored document path is invalid")
    return resolved


async def store_upload(
    upload: UploadSource,
    *,
    tenant_id: UUID,
    document_id: UUID,
    settings: Settings,
) -> StoredUpload:
    filename = sanitize_filename(upload.filename)
    extension = _extension(filename)
    _validate_declared_media_type(extension, upload.content_type)

    content = bytearray()
    while block := await upload.read(READ_BLOCK_BYTES):
        content.extend(block)
        if len(content) > settings.document_max_upload_bytes:
            raise DocumentUploadError(413, "The uploaded document exceeds the size limit")
    data = bytes(content)
    _validate_file_signature(extension, data)

    storage_key = f"{tenant_id}/{document_id}/{uuid4().hex}{extension}"
    path = resolve_storage_path(settings.document_storage_root, storage_key)

    def write_file() -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("xb") as destination:
            destination.write(data)

    try:
        await asyncio.to_thread(write_file)
    except OSError as exc:
        try:
            await asyncio.to_thread(path.unlink, missing_ok=True)
        except OSError:
            pass
        raise DocumentUploadError(500, "The document could not be stored") from exc

    return StoredUpload(
        original_filename=filename,
        storage_key=storage_key,
        media_type=_MEDIA_TYPES[extension],
        size_bytes=len(data),
        file_sha256=hashlib.sha256(data).hexdigest(),
    )


async def remove_stored_upload(storage_root: Path, storage_key: str) -> None:
    try:
        path = resolve_storage_path(storage_root, storage_key)
    except DocumentProcessingError:
        return
    try:
        await asyncio.to_thread(path.unlink, missing_ok=True)
    except OSError:
        return


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [re.sub(r"[\t ]+", " ", line).strip() for line in text.splitlines()]
    normalized: list[str] = []
    previous_blank = False
    for line in lines:
        is_blank = not line
        if is_blank and previous_blank:
            continue
        normalized.append(line)
        previous_blank = is_blank
    return "\n".join(normalized).strip()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader  # type: ignore[import-not-found]
    except ImportError as exc:
        raise DocumentProcessingError("PDF processing support is not installed") from exc
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise DocumentProcessingError("Encrypted PDF files are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentProcessingError("The PDF contains too many pages")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("The PDF could not be read") from exc


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document  # type: ignore[import-not-found]
    except ImportError as exc:
        raise DocumentProcessingError("DOCX processing support is not installed") from exc
    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise DocumentProcessingError("The DOCX could not be read") from exc


def _extract_html(data: bytes) -> str:
    try:
        from bs4 import BeautifulSoup  # type: ignore[import-not-found]
    except ImportError as exc:
        raise DocumentProcessingError("HTML processing support is not installed") from exc
    soup = BeautifulSoup(data.decode("utf-8-sig"), "html.parser")
    for unwanted in soup(["script", "style", "noscript"]):
        unwanted.decompose()
    return str(soup.get_text("\n"))


def extract_document_text(
    data: bytes,
    *,
    filename: str,
    maximum_characters: int,
) -> str:
    extension = _extension(filename)
    _validate_file_signature(extension, data)
    if extension in {".txt", ".md"}:
        raw_text = data.decode("utf-8-sig")
    elif extension in {".html", ".htm"}:
        raw_text = _extract_html(data)
    elif extension == ".pdf":
        raw_text = _extract_pdf(data)
    else:
        raw_text = _extract_docx(data)

    text = _normalize_text(raw_text)
    if not text:
        raise DocumentProcessingError("No readable text was found in the document")
    if len(text) > maximum_characters:
        raise DocumentProcessingError("The extracted document text is too large")
    return text


def chunk_text(text: str, *, maximum_characters: int, overlap: int) -> list[str]:
    if maximum_characters < 1 or overlap < 0 or overlap >= maximum_characters:
        raise ValueError("Chunk size must be positive and overlap must be smaller than it")
    if len(text) <= maximum_characters:
        return [text]

    chunks: list[str] = []
    start = 0
    while start < len(text):
        target = min(start + maximum_characters, len(text))
        end = target
        if target < len(text):
            split = max(text.rfind("\n", start + maximum_characters // 2, target), 0)
            if split <= start:
                split = text.rfind(" ", start + maximum_characters // 2, target)
            if split > start:
                end = split
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        next_start = max(0, end - overlap)
        while next_start < end and text[next_start].isspace():
            next_start += 1
        start = next_start if next_start > start else end
    return chunks


class InProcessDocumentQueue:
    def __init__(
        self,
        session_factory: async_sessionmaker[AsyncSession],
        *,
        settings: Settings,
        embedding_provider: EmbeddingProvider | None,
    ) -> None:
        self._session_factory = session_factory
        self._settings = settings
        self._embedding_provider = embedding_provider

    async def process(self, document_id: UUID) -> None:
        async with self._session_factory() as session:
            document = await self._claim_document(session, document_id)
            if document is None:
                return
            try:
                await self._process_claimed_document(session, document)
            except Exception as exc:
                await session.rollback()
                await self._mark_failed(session, document_id, exc)

    async def _claim_document(
        self, session: AsyncSession, document_id: UUID
    ) -> KnowledgeDocument | None:
        document = (
            await session.execute(
                select(KnowledgeDocument)
                .where(KnowledgeDocument.id == document_id)
                .with_for_update()
            )
        ).scalar_one_or_none()
        if document is None or document.processing_status not in {
            DocumentProcessingStatus.UPLOADED,
            DocumentProcessingStatus.FAILED,
        }:
            return None
        document.processing_status = DocumentProcessingStatus.PROCESSING
        document.processing_progress = 5
        document.processing_stage = "reading_file"
        document.processing_eta_seconds = 15
        document.processing_error = None
        await session.commit()
        return document

    @staticmethod
    async def _record_progress(
        session: AsyncSession,
        document: KnowledgeDocument,
        *,
        progress: int,
        stage: str,
        eta_seconds: int,
    ) -> None:
        document.processing_progress = max(0, min(99, progress))
        document.processing_stage = stage
        document.processing_eta_seconds = max(0, eta_seconds)
        await session.commit()

    async def _process_claimed_document(
        self, session: AsyncSession, document: KnowledgeDocument
    ) -> None:
        if not document.storage_key or not document.original_filename:
            raise DocumentProcessingError("Document storage metadata is missing")
        path = resolve_storage_path(self._settings.document_storage_root, document.storage_key)
        try:
            data = await asyncio.to_thread(path.read_bytes)
        except OSError as exc:
            raise DocumentProcessingError("The stored document is unavailable") from exc
        if hashlib.sha256(data).hexdigest() != document.file_sha256:
            raise DocumentProcessingError("The stored document failed its integrity check")

        await self._record_progress(
            session,
            document,
            progress=20,
            stage="extracting_text",
            eta_seconds=12,
        )
        text = await asyncio.to_thread(
            extract_document_text,
            data,
            filename=document.original_filename,
            maximum_characters=self._settings.document_max_extracted_characters,
        )
        await self._record_progress(
            session,
            document,
            progress=45,
            stage="chunking",
            eta_seconds=8,
        )
        chunks = chunk_text(
            text,
            maximum_characters=self._settings.document_chunk_characters,
            overlap=self._settings.document_chunk_overlap,
        )
        embedding_batches = max(1, (len(chunks) + EMBEDDING_BATCH_SIZE - 1) // EMBEDDING_BATCH_SIZE)
        await self._record_progress(
            session,
            document,
            progress=55,
            stage="embedding",
            eta_seconds=embedding_batches * 4 + 3,
        )
        embeddings, embedding_status = await self._embed_chunks(session, document, chunks)
        await self._record_progress(
            session,
            document,
            progress=92,
            stage="saving",
            eta_seconds=2,
        )
        current_version = await session.scalar(
            select(func.coalesce(func.max(KnowledgeVersion.version), 0)).where(
                KnowledgeVersion.document_id == document.id
            )
        )
        next_version = (current_version or 0) + 1
        version = KnowledgeVersion(
            document_id=document.id,
            version=next_version,
            content=text,
            content_sha256=hashlib.sha256(text.encode("utf-8")).hexdigest(),
            source=document.original_filename,
            embedding_status=embedding_status,
            metadata_json={
                "filename": document.original_filename,
                "media_type": document.media_type,
                "chunk_count": len(chunks),
            },
        )
        session.add(version)
        await session.flush()
        session.add_all(
            [
                KnowledgeChunk(
                    version_id=version.id,
                    tenant_id=document.tenant_id,
                    property_id=document.property_id,
                    chunk_index=index,
                    content=chunk,
                    token_count=max(1, len(chunk) // 4),
                    embedding=embeddings[index],
                    metadata_json={
                        "filename": document.original_filename,
                        "chunk_index": index,
                    },
                )
                for index, chunk in enumerate(chunks)
            ]
        )
        document.status = KnowledgeStatus.PENDING_REVIEW
        document.processing_status = DocumentProcessingStatus.NEEDS_REVIEW
        document.processing_progress = 100
        document.processing_stage = "complete"
        document.processing_eta_seconds = 0
        document.processed_at = datetime.now(timezone.utc)
        document.processing_error = None
        session.add(
            AuditEvent(
                tenant_id=document.tenant_id,
                actor_user_id=document.uploaded_by_id,
                event_type="document.processed",
                entity_type="knowledge_document",
                entity_id=document.id,
                details={
                    "chunk_count": len(chunks),
                    "embedding_status": embedding_status,
                },
            )
        )
        await session.commit()

    async def _embed_chunks(
        self,
        session: AsyncSession,
        document: KnowledgeDocument,
        chunks: list[str],
    ) -> tuple[list[list[float] | None], str]:
        if self._embedding_provider is None:
            await self._record_progress(
                session,
                document,
                progress=88,
                stage="embedding_deferred",
                eta_seconds=3,
            )
            return [None] * len(chunks), "pending"
        try:
            embeddings: list[list[float] | None] = []
            for start in range(0, len(chunks), EMBEDDING_BATCH_SIZE):
                batch = chunks[start : start + EMBEDDING_BATCH_SIZE]
                embeddings.extend(await self._embedding_provider.embed_documents(batch))
                completed = min(start + len(batch), len(chunks))
                progress = 55 + round(33 * completed / len(chunks))
                remaining_batches = (
                    len(chunks) - completed + EMBEDDING_BATCH_SIZE - 1
                ) // EMBEDDING_BATCH_SIZE
                await self._record_progress(
                    session,
                    document,
                    progress=progress,
                    stage="embedding",
                    eta_seconds=remaining_batches * 4 + 3,
                )
        except EmbeddingProviderError:
            await self._record_progress(
                session,
                document,
                progress=88,
                stage="embedding_deferred",
                eta_seconds=3,
            )
            return [None] * len(chunks), "pending"
        return embeddings, "ready"

    async def _mark_failed(self, session: AsyncSession, document_id: UUID, exc: Exception) -> None:
        document = await session.get(KnowledgeDocument, document_id)
        if document is None:
            return
        detail = (
            str(exc) if isinstance(exc, DocumentProcessingError) else "Document processing failed"
        )
        document.processing_status = DocumentProcessingStatus.FAILED
        document.processing_stage = "failed"
        document.processing_eta_seconds = None
        document.processing_error = detail[:1000]
        document.processed_at = datetime.now(timezone.utc)
        session.add(
            AuditEvent(
                tenant_id=document.tenant_id,
                actor_user_id=document.uploaded_by_id,
                event_type="document.processing_failed",
                entity_type="knowledge_document",
                entity_id=document.id,
                details={"reason": detail[:200]},
            )
        )
        await session.commit()
